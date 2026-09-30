"""Secure Excel review projection for canonical JSON test plans."""
import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook

from agent.guardrails import check_step
from agent.models import Step, Target, TestPlan

MAX_WORKBOOK_BYTES = 10 * 1024 * 1024
MAX_STEPS = 10_000
DANGEROUS_PREFIXES = ("=", "+", "-", "@")
REVIEW_STATUSES = {"DRAFT", "APPROVED", "REJECTED", "NEEDS_CHANGE"}


def _safe_cell(value):
    if isinstance(value, str) and value.startswith(DANGEROUS_PREFIXES):
        return "'" + value
    return value


def _reject_unsafe(value, secrets: Iterable[str]) -> None:
    if not isinstance(value, str):
        return
    if value.startswith(DANGEROUS_PREFIXES):
        raise ValueError("Spreadsheet formulas and formula-like values are not allowed")
    lower = value.lower()
    if lower.startswith(("bearer ", "basic ")) or "sk-" in lower:
        raise ValueError("Credential-like value found in workbook")
    if any(secret and secret in value for secret in secrets):
        raise ValueError("A configured secret was found in workbook")


def export_plan(plan: TestPlan, path: Path, secrets: Iterable[str] = ()) -> Path:
    if path.suffix.lower() != ".xlsx":
        raise ValueError("Only macro-free .xlsx workbooks are supported")
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    for key, value in (
        ("plan_id", plan.id), ("version", plan.version), ("story_key", plan.story_key),
        ("story_title", plan.story_title), ("created_at", plan.created_at),
        ("generator", plan.generator),
    ):
        summary.append([key, _safe_cell(value)])

    cases = workbook.create_sheet("Test Cases")
    case_headers = [
        "case_id", "story_key", "ac_ids", "title", "rationale", "status",
        "reviewer", "review_comment", "reviewed_at", "feature_ids",
        "figma_frames", "figma_node_ids",
    ]
    cases.append(case_headers)
    for case in plan.cases:
        row = [
            case.id, case.story_key, ",".join(case.ac_ids), case.title, case.rationale,
            case.status, case.reviewer, case.review_comment, case.reviewed_at or "",
            ",".join(case.feature_ids), ",".join(case.figma_frames), ",".join(case.figma_node_ids),
        ]
        for value in row:
            _reject_unsafe(value, secrets)
        cases.append([_safe_cell(v) for v in row])

    steps = workbook.create_sheet("Steps")
    step_headers = [
        "case_id", "index", "action", "target_role", "target_name", "target_label",
        "target_text", "target_testid", "value", "values_json", "path", "contains",
        "expected_text", "name", "frame",
    ]
    steps.append(step_headers)
    for case in plan.cases:
        for index, step in enumerate(case.steps, 1):
            target = step.target or Target()
            row = [
                case.id, index, step.action, target.role, target.name, target.label,
                target.text, target.testid, step.value, json.dumps(step.values) if step.values is not None else "",
                step.path, step.contains, step.text, step.name, step.frame,
            ]
            for value in row:
                _reject_unsafe(value, secrets)
            steps.append([_safe_cell(v) for v in row])

    coverage = workbook.create_sheet("Coverage")
    coverage.append(["ac_id", "case_ids"])
    acs = sorted({ac for case in plan.cases for ac in case.ac_ids})
    for ac in acs:
        coverage.append([ac, ",".join(case.id for case in plan.cases if ac in case.ac_ids)])

    sources = workbook.create_sheet("Sources")
    sources.append(["citation_id", "source", "chunk_id", "score", "excerpt"])
    for citation in plan.citations:
        sources.append([
            citation.id, citation.source, citation.chunk_id, citation.score,
            _safe_cell(citation.excerpt),
        ])

    review = workbook.create_sheet("Human Review")
    review.append(["case_id", "status", "reviewer", "comment"])
    for case in plan.cases:
        review.append([case.id, case.status, case.reviewer, _safe_cell(case.review_comment)])

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def _rows(sheet) -> list[dict]:
    values = list(sheet.iter_rows(values_only=True))
    if not values:
        return []
    headers = [str(value or "").strip() for value in values[0]]
    if len(headers) != len(set(headers)):
        raise ValueError(f"Duplicate columns in {sheet.title}")
    return [dict(zip(headers, row)) for row in values[1:] if any(v is not None for v in row)]


def import_review(
    path: Path,
    canonical: TestPlan,
    *,
    base_url: str,
    secrets: Iterable[str] = (),
    valid_ac_ids: set[str] | None = None,
) -> TestPlan:
    if path.suffix.lower() != ".xlsx":
        raise ValueError("Only macro-free .xlsx workbooks are supported")
    if path.stat().st_size > MAX_WORKBOOK_BYTES:
        raise ValueError("Workbook is too large")
    workbook = load_workbook(path, data_only=False, read_only=True)
    required = {"Summary", "Test Cases", "Steps", "Human Review"}
    if not required.issubset(workbook.sheetnames):
        raise ValueError(f"Workbook is missing sheets: {sorted(required - set(workbook.sheetnames))}")

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            for value in row:
                _reject_unsafe(value, secrets)

    summary = {row[0].value: row[1].value for row in workbook["Summary"].iter_rows(min_row=1, max_col=2)}
    if summary.get("plan_id") != canonical.id or int(summary.get("version", 0)) != canonical.version:
        raise ValueError("Workbook does not match the canonical plan id/version")

    by_id = {case.id: case.model_copy(deep=True) for case in canonical.cases}
    case_rows = _rows(workbook["Test Cases"])
    allowed_case_columns = {
        "case_id", "story_key", "ac_ids", "title", "rationale", "status", "reviewer",
        "review_comment", "reviewed_at", "feature_ids", "figma_frames", "figma_node_ids",
    }
    if case_rows and set(case_rows[0]) - allowed_case_columns:
        raise ValueError("Unknown Test Cases columns")
    for row in case_rows:
        case_id = str(row.get("case_id") or "")
        if case_id not in by_id:
            raise ValueError(f"Unknown case id {case_id}")
        ac_ids = [v.strip() for v in str(row.get("ac_ids") or "").split(",") if v.strip()]
        if valid_ac_ids is not None and not set(ac_ids).issubset(valid_ac_ids):
            raise ValueError(f"{case_id} cites unknown acceptance criteria")
        case = by_id[case_id]
        case.ac_ids = ac_ids
        case.title = str(row.get("title") or "")
        case.rationale = str(row.get("rationale") or "")

    step_rows = _rows(workbook["Steps"])
    if len(step_rows) > MAX_STEPS:
        raise ValueError("Workbook contains too many steps")
    grouped: dict[str, list[tuple[int, Step]]] = {case_id: [] for case_id in by_id}
    for row in step_rows:
        case_id = str(row.get("case_id") or "")
        if case_id not in by_id:
            raise ValueError(f"Unknown case id {case_id}")
        target_values = {
            "role": row.get("target_role"), "name": row.get("target_name"),
            "label": row.get("target_label"), "text": row.get("target_text"),
            "testid": row.get("target_testid"),
        }
        target_values = {k: str(v) for k, v in target_values.items() if v not in (None, "")}
        values = row.get("values_json")
        step = Step(
            action=str(row.get("action") or ""),
            target=Target(**target_values) if target_values else None,
            value=str(row["value"]) if row.get("value") not in (None, "") else None,
            values=json.loads(values) if values not in (None, "") else None,
            path=str(row["path"]) if row.get("path") not in (None, "") else None,
            contains=str(row["contains"]) if row.get("contains") not in (None, "") else None,
            text=str(row["expected_text"]) if row.get("expected_text") not in (None, "") else None,
            name=str(row["name"]) if row.get("name") not in (None, "") else None,
            frame=str(row["frame"]) if row.get("frame") not in (None, "") else None,
        )
        refusal = check_step(step, base_url)
        if refusal:
            raise ValueError(f"{case_id} step refused: {refusal}")
        grouped[case_id].append((int(row.get("index") or 0), step))
    for case_id, indexed in grouped.items():
        if indexed:
            indexes = [i for i, _ in indexed]
            if indexes != list(range(1, len(indexes) + 1)):
                raise ValueError(f"{case_id} step indexes must be contiguous and ordered")
            by_id[case_id].steps = [step for _, step in indexed]

    review_rows = _rows(workbook["Human Review"])
    for row in review_rows:
        case_id = str(row.get("case_id") or "")
        if case_id not in by_id:
            raise ValueError(f"Unknown case id {case_id}")
        status = str(row.get("status") or "DRAFT").upper()
        if status not in REVIEW_STATUSES:
            raise ValueError(f"Invalid review status {status}")
        case = by_id[case_id]
        case.status = status
        case.reviewer = str(row.get("reviewer") or "")
        case.review_comment = str(row.get("comment") or "")
        case.reviewed_at = datetime.now(timezone.utc).isoformat() if status != "DRAFT" else None

    reviewed = canonical.model_copy(deep=True)
    reviewed.cases = [by_id[case.id] for case in canonical.cases]
    return reviewed


def append_execution_results(
    path: Path,
    story_results: list,
    evaluation: dict,
    secrets: Iterable[str] = (),
    performance: dict | None = None,
) -> Path:
    """Add replaceable execution and evaluation sheets to a review workbook."""
    if not path.is_file():
        return path
    workbook = load_workbook(path)
    for name in ("Results", "Evaluation", "Performance"):
        if name in workbook.sheetnames:
            del workbook[name]
    results_sheet = workbook.create_sheet("Results")
    results_sheet.append(["story_key", "ac_id", "label", "rationale", "evidence"])
    for result in story_results:
        payload = result.model_dump() if hasattr(result, "model_dump") else result
        for verdict in payload.get("ac_verdicts", []):
            row = [
                payload.get("story_key"), verdict.get("ac_id"), verdict.get("label"),
                verdict.get("rationale"), "\n".join(verdict.get("evidence", [])),
            ]
            for value in row:
                _reject_unsafe(value, secrets)
            results_sheet.append([_safe_cell(value) for value in row])
    eval_sheet = workbook.create_sheet("Evaluation")
    eval_sheet.append(["check", "passed", "details"])
    for name, check in evaluation.get("checks", {}).items():
        details = json.dumps({k: v for k, v in check.items() if k != "passed"}, default=str)
        _reject_unsafe(details, secrets)
        eval_sheet.append([name, bool(check.get("passed")), _safe_cell(details)])
    eval_sheet.append(["overall", bool(evaluation.get("passed")), ""])
    if performance:
        perf_sheet = workbook.create_sheet("Performance")
        perf_sheet.append([
            "profile", "scope", "status", "metric", "count",
            "median", "p90", "p95", "p99", "cv", "budget",
        ])
        for summary in performance.get("summaries", []):
            budgets = {
                item.get("metric"): item for item in summary.get("budgets", [])
            }
            for metric, stats in summary.get("statistics", {}).items():
                budget = budgets.get(metric, {})
                detail = budget.get("reason", "")
                _reject_unsafe(detail, secrets)
                perf_sheet.append([
                    summary.get("profile_id"), summary.get("scope"), summary.get("status"),
                    metric, stats.get("count"), stats.get("median"), stats.get("p90"),
                    stats.get("p95"), stats.get("p99"), stats.get("cv"), _safe_cell(detail),
                ])
    workbook.save(path)
    return path
