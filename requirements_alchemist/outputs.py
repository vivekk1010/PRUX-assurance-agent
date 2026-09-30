"""Human-review Excel workflow and guarded Jira Cloud publication."""
from __future__ import annotations

import json
import os
from datetime import datetime, timezone
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import httpx
from openpyxl import Workbook, load_workbook

from agent.models import AcceptanceCriterion as AssuranceCriterion
from agent.models import Story as AssuranceStory
from requirements_alchemist.config import Settings
from requirements_alchemist.models import GeneratedStory, PublishResult, StoryPackage

DANGEROUS_PREFIXES = ("=", "+", "-", "@")
REVIEW_STATUSES = {"DRAFT", "APPROVED", "REJECTED", "NEEDS_CHANGE"}
MAX_WORKBOOK_BYTES = 15 * 1024 * 1024


def _safe(value):
    return "'" + value if isinstance(value, str) and value.startswith(DANGEROUS_PREFIXES) else value


def _reject_unsafe(
    value, secrets: list[str], *, reject_formulas: bool = True
) -> None:
    if not isinstance(value, str):
        return
    if reject_formulas and value.startswith(DANGEROUS_PREFIXES):
        raise ValueError("Formula-like workbook values are not accepted")
    lowered = value.lower()
    if lowered.startswith(("bearer ", "basic ")) or "sk-" in lowered or any(
        secret and secret in value for secret in secrets
    ):
        raise ValueError("Credential-like content found in workbook")


def _append_export(sheet, row: list, secrets: list[str]) -> None:
    for value in row:
        _reject_unsafe(value, secrets, reject_formulas=False)
    sheet.append([_safe(value) for value in row])


def export_excel(package: StoryPackage, path: Path, secrets: list[str] | None = None) -> Path:
    if path.suffix.lower() != ".xlsx":
        raise ValueError("Story review output must be a macro-free .xlsx file")
    secrets = secrets or []
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Package"
    for row in (
        ("package_id", package.id),
        ("title", package.title),
        ("objective", package.objective),
        ("generated_at", package.generated_at),
        ("generator", package.generator),
        ("instructions", "Review Stories and set Human Review status to APPROVED before Jira push."),
    ):
        _append_export(summary, list(row), secrets)

    stories = workbook.create_sheet("Stories")
    stories.append(
        [
            "local_id",
            "epic",
            "title",
            "narrative",
            "business_value",
            "description",
            "priority",
            "story_points",
            "confidence",
            "scope_in",
            "scope_out",
            "dependencies",
            "assumptions",
            "risks",
            "open_questions",
            "labels",
            "jira_key",
        ]
    )
    for story in package.stories:
        row = [
            story.local_id,
            story.epic,
            story.title,
            story.narrative,
            story.business_value,
            story.description,
            story.priority,
            story.story_points,
            story.confidence,
            "\n".join(story.scope_in),
            "\n".join(story.scope_out),
            "\n".join(story.dependencies),
            "\n".join(story.assumptions),
            "\n".join(story.risks),
            "\n".join(story.open_questions),
            ",".join(story.labels),
            story.jira_key or "",
        ]
        _append_export(stories, row, secrets)

    criteria = workbook.create_sheet("Acceptance Criteria")
    criteria.append(
        ["local_id", "ac_id", "title", "given", "when", "then", "negative", "source_ids"]
    )
    for story in package.stories:
        for criterion in story.acceptance_criteria:
            _append_export(
                criteria,
                [
                    story.local_id,
                    criterion.id,
                    criterion.title,
                    "\n".join(criterion.given),
                    criterion.when,
                    "\n".join(criterion.then),
                    criterion.negative,
                    ",".join(criterion.source_ids),
                ],
                secrets,
            )

    nfrs = workbook.create_sheet("Non-functional")
    nfrs.append(["local_id", "category", "requirement", "measure", "source_ids"])
    for story in package.stories:
        for nfr in story.non_functional_requirements:
            _append_export(
                nfrs,
                [
                    story.local_id,
                    nfr.category,
                    nfr.requirement,
                    nfr.measure,
                    ",".join(nfr.source_ids),
                ],
                secrets,
            )

    review = workbook.create_sheet("Human Review")
    review.append(["local_id", "status", "reviewer", "comment"])
    for story in package.stories:
        _append_export(
            review,
            [
                story.local_id,
                story.review_status,
                story.reviewer,
                story.review_comment,
            ],
            secrets,
        )

    sources = workbook.create_sheet("Sources")
    sources.append(["local_id", "source_id", "title", "locator", "excerpt"])
    for story in package.stories:
        for citation in story.citations:
            _append_export(
                sources,
                [
                    story.local_id,
                    citation.source_id,
                    citation.source_title,
                    citation.locator,
                    citation.excerpt,
                ],
                secrets,
            )
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def import_excel_review(
    path: Path, package: StoryPackage, secrets: list[str] | None = None
) -> StoryPackage:
    if path.suffix.lower() != ".xlsx" or path.stat().st_size > MAX_WORKBOOK_BYTES:
        raise ValueError("Invalid or oversized review workbook")
    workbook = load_workbook(path, data_only=False, read_only=True)
    required = {"Package", "Stories", "Acceptance Criteria", "Human Review"}
    if not required.issubset(workbook.sheetnames):
        raise ValueError(f"Workbook is missing sheets: {sorted(required - set(workbook.sheetnames))}")
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            for value in row:
                _reject_unsafe(value, secrets or [])
    metadata = {
        row[0].value: row[1].value
        for row in workbook["Package"].iter_rows(min_row=1, max_col=2)
    }
    if metadata.get("package_id") != package.id:
        raise ValueError("Workbook package id does not match the canonical JSON")

    values = list(workbook["Human Review"].iter_rows(values_only=True))
    headers = [str(value or "") for value in values[0]]
    if headers != ["local_id", "status", "reviewer", "comment"]:
        raise ValueError("Human Review columns were changed")
    by_id = {story.local_id: story.model_copy(deep=True) for story in package.stories}
    for row in values[1:]:
        item = dict(zip(headers, row))
        local_id = str(item.get("local_id") or "")
        if local_id not in by_id:
            raise ValueError(f"Unknown story id in review: {local_id}")
        status = str(item.get("status") or "DRAFT").upper()
        if status not in REVIEW_STATUSES:
            raise ValueError(f"Invalid review status: {status}")
        by_id[local_id].review_status = status
        by_id[local_id].reviewer = str(item.get("reviewer") or "")
        by_id[local_id].review_comment = str(item.get("comment") or "")
    reviewed = package.model_copy(deep=True)
    reviewed.stories = [by_id[story.local_id] for story in package.stories]
    return reviewed


def _paragraph(text: str) -> dict:
    return {
        "type": "paragraph",
        "content": [{"type": "text", "text": text[:30_000]}],
    }


def jira_description(story: GeneratedStory) -> dict:
    blocks = [
        _paragraph(story.narrative),
        _paragraph(f"Business value: {story.business_value}"),
        _paragraph(story.description),
        _paragraph("Acceptance criteria"),
    ]
    for criterion in story.acceptance_criteria:
        given = " AND ".join(criterion.given)
        then = " AND ".join(criterion.then)
        blocks.append(
            _paragraph(
                f"{criterion.id} — {criterion.title}\nGIVEN {given}\n"
                f"WHEN {criterion.when}\nTHEN {then}"
            )
        )
    if story.open_questions:
        blocks.append(_paragraph("Open questions\n" + "\n".join(story.open_questions)))
    if story.citations:
        blocks.append(
            _paragraph(
                "Source traceability\n"
                + "\n".join(
                    f"{citation.source_id}: {citation.source_title}"
                    for citation in story.citations
                )
            )
        )
    return {"type": "doc", "version": 1, "content": blocks}


class JiraPublisher:
    def __init__(self, settings: Settings):
        self.settings = settings

    def _auth(self) -> tuple[str, str]:
        email = os.getenv(self.settings.atlassian_email_env, "")
        token = os.getenv(self.settings.atlassian_token_env, "")
        if not email or not token:
            raise ValueError("Atlassian credentials are not configured")
        return email, token

    def publish(self, package: StoryPackage) -> list[PublishResult]:
        if not self.settings.jira_push_enabled:
            raise PermissionError("Jira publishing is disabled by configuration")
        if not self.settings.atlassian_base_url or not self.settings.jira_project_key:
            raise ValueError("Atlassian base URL and Jira project key are required")
        results: list[PublishResult] = []
        endpoint = f"{self.settings.atlassian_base_url.rstrip('/')}/rest/api/3/issue"
        for story in package.stories:
            if story.jira_key:
                results.append(
                    PublishResult(
                        local_id=story.local_id,
                        status="SKIPPED",
                        jira_key=story.jira_key,
                        reason="Already published",
                    )
                )
                continue
            if self.settings.require_human_approval and story.review_status != "APPROVED":
                results.append(
                    PublishResult(
                        local_id=story.local_id,
                        status="SKIPPED",
                        reason="Human approval is required",
                    )
                )
                continue
            payload = {
                "fields": {
                    "project": {"key": self.settings.jira_project_key},
                    "issuetype": {"name": self.settings.jira_issue_type},
                    "summary": story.title[:255],
                    "description": jira_description(story),
                    "labels": sorted(set(story.labels + ["requirements-alchemist"])),
                },
                "properties": [
                    {
                        "key": "requirements-alchemist",
                        "value": {
                            "package_id": package.id,
                            "local_id": story.local_id,
                            "published_at": datetime.now(timezone.utc).isoformat(),
                        },
                    }
                ],
            }
            try:
                response = httpx.post(endpoint, auth=self._auth(), json=payload, timeout=30)
                response.raise_for_status()
                key = response.json()["key"]
                story.jira_key = key
                results.append(
                    PublishResult(
                        local_id=story.local_id,
                        status="CREATED",
                        jira_key=key,
                        url=f"{self.settings.atlassian_base_url.rstrip('/')}/browse/{key}",
                    )
                )
            except Exception as exc:
                results.append(
                    PublishResult(
                        local_id=story.local_id,
                        status="FAILED",
                        reason=f"{type(exc).__name__}: {exc}",
                    )
                )
        return results


def save_package(package: StoryPackage, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(package.model_dump_json(indent=2), encoding="utf-8")
    return path


def load_package(path: Path) -> StoryPackage:
    return StoryPackage.model_validate_json(path.read_text(encoding="utf-8"))


def export_assurance_bundle(
    package: StoryPackage, path: Path, *, approved_only: bool = True
) -> Path:
    """Project approved generated stories into the existing assurance-agent contract."""
    selected = [
        story
        for story in package.stories
        if not approved_only or story.review_status == "APPROVED"
    ]
    if not selected:
        raise ValueError("No approved stories are available for assurance export")
    path.parent.mkdir(parents=True, exist_ok=True)
    with ZipFile(path, "w", compression=ZIP_DEFLATED) as archive:
        for story in selected:
            projected = AssuranceStory(
                key=story.jira_key or story.local_id,
                title=story.title,
                status=story.review_status,
                priority=story.priority,
                actor=story.persona,
                description="\n\n".join(
                    [
                        story.narrative,
                        story.business_value,
                        story.description,
                        "Open questions:\n" + "\n".join(story.open_questions)
                        if story.open_questions
                        else "",
                    ]
                ).strip(),
                acceptance_criteria=[
                    AssuranceCriterion(
                        id=criterion.id,
                        text=(
                            f"{criterion.title}: GIVEN {' AND '.join(criterion.given)} "
                            f"WHEN {criterion.when} THEN {' AND '.join(criterion.then)}"
                        ),
                    )
                    for criterion in story.acceptance_criteria
                ],
                business_rules=(
                    story.scope_in
                    + story.data_requirements
                    + [
                        f"{item.category}: {item.requirement} ({item.measure})"
                        for item in story.non_functional_requirements
                    ]
                ),
                labels=story.labels,
            )
            archive.writestr(
                f"stories/{projected.key}.json",
                projected.model_dump_json(indent=2),
            )
    return path
