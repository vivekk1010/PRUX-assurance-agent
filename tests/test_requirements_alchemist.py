import json

import pytest

from requirements_alchemist.app import create_app
from requirements_alchemist.config import ROOT, Settings
from requirements_alchemist.generation import RequirementsLLM
from requirements_alchemist.models import (
    AcceptanceCriterion,
    GeneratedStory,
    GenerationInput,
    StoryPackage,
)
from requirements_alchemist.outputs import (
    JiraPublisher,
    export_assurance_bundle,
    export_excel,
    import_excel_review,
    jira_description,
)
from requirements_alchemist.retrieval import ReferenceCorpus, RetrievedChunk
from requirements_alchemist.sources import SourceLoader


def _story(status="DRAFT"):
    return GeneratedStory(
        local_id="STORY-001",
        epic="Identity",
        title="Sign in with an approved account",
        persona="registered user",
        capability="sign in",
        benefit="I can access my workspace",
        narrative="",
        business_value="Protects and enables account access.",
        description="Authenticate an active account.",
        acceptance_criteria=[
            AcceptanceCriterion(
                id="AC-01",
                title="Valid credentials",
                given=["an active account exists"],
                when="valid credentials are submitted",
                then=["the workspace opens"],
                source_ids=["prd-1"],
            )
        ],
        review_status=status,
        reviewer="Vivek" if status == "APPROVED" else "",
    )


def _package(status="DRAFT"):
    return StoryPackage(
        id="pkg-test",
        title="Identity backlog",
        objective="Enable secure access",
        source_ids=["prd-1"],
        stories=[_story(status)],
    )


def test_reference_corpus_retrieves_licensed_story_pattern(tmp_path):
    (tmp_path / "samples.json").write_text(
        json.dumps(
            [
                {
                    "id": "reference-access",
                    "title": "Accessible authentication",
                    "text": "As a keyboard user I can sign in with visible focus and error recovery.",
                    "license": "CC0-1.0",
                    "source_url": "https://example.test/reference",
                }
            ]
        ),
        encoding="utf-8",
    )
    corpus = ReferenceCorpus(tmp_path)
    corpus.rebuild()

    hits = corpus.search("keyboard sign in focus")

    assert hits[0].source_id == "reference-access"
    assert hits[0].metadata["license"] == "CC0-1.0"


def test_local_ollama_profile_is_keyless_and_write_safe(monkeypatch):
    for name in ("RA_LLM_PROVIDER", "RA_LLM_MODEL", "RA_LLM_BASE_URL"):
        monkeypatch.delenv(name, raising=False)
    settings = Settings.load(ROOT / "config" / "requirements-alchemist.ollama.json")

    assert settings.llm_provider == "openai-compatible"
    assert settings.llm_model == "qwen2.5:7b"
    assert settings.llm_base_url == "http://127.0.0.1:11434/v1"
    assert settings.jira_push_enabled is False
    assert settings.require_human_approval is True


def test_excel_review_only_updates_review_fields(tmp_path):
    package = _package()
    path = export_excel(package, tmp_path / "review.xlsx")
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    review = workbook["Human Review"]
    review["B2"] = "APPROVED"
    review["C2"] = "Vivek"
    review["D2"] = "Evidence and scope reviewed"
    workbook["Stories"]["C2"] = "Untrusted spreadsheet title edit"
    workbook.save(path)

    reviewed = import_excel_review(path, package)

    assert reviewed.stories[0].review_status == "APPROVED"
    assert reviewed.stories[0].reviewer == "Vivek"
    assert reviewed.stories[0].title == package.stories[0].title


def test_excel_export_neutralizes_formula_like_story_text(tmp_path):
    from openpyxl import load_workbook

    package = _package()
    package.stories[0].title = "=HYPERLINK(\"https://example.test\")"
    path = export_excel(package, tmp_path / "safe.xlsx")

    workbook = load_workbook(path, data_only=False, read_only=True)
    assert workbook["Stories"]["C2"].value.startswith("'=")


def test_jira_description_contains_traceable_acceptance_criteria():
    document = jira_description(_story("APPROVED"))
    encoded = json.dumps(document)

    assert "AC-01" in encoded
    assert "GIVEN an active account exists" in encoded
    assert document["type"] == "doc"


def test_approved_story_exports_to_assurance_contract(tmp_path):
    from zipfile import ZipFile

    path = export_assurance_bundle(
        _package("APPROVED"), tmp_path / "assurance.zip"
    )

    with ZipFile(path) as archive:
        payload = json.loads(archive.read("stories/STORY-001.json"))
    assert payload["actor"] == "registered user"
    assert payload["acceptance_criteria"][0]["id"] == "AC-01"


def test_jira_publisher_skips_unapproved_story(tmp_path):
    settings = Settings(
        jira_push_enabled=True,
        require_human_approval=True,
        atlassian_base_url="https://example.atlassian.net",
        jira_project_key="PROJ",
        reference_dir=tmp_path,
        workspace_dir=tmp_path,
        output_dir=tmp_path,
    )

    result = JiraPublisher(settings).publish(_package("DRAFT"))

    assert result[0].status == "SKIPPED"
    assert result[0].reason == "Human approval is required"


def test_source_loader_rejects_non_https_url(tmp_path):
    loader = SourceLoader(
        Settings(
            reference_dir=tmp_path,
            workspace_dir=tmp_path,
            output_dir=tmp_path,
        )
    )

    with pytest.raises(ValueError, match="HTTPS"):
        loader.load_url("http://127.0.0.1/private")


def test_source_loader_fails_closed_for_unlisted_host(tmp_path):
    loader = SourceLoader(
        Settings(
            allowed_source_hosts=(),
            reference_dir=tmp_path,
            workspace_dir=tmp_path,
            output_dir=tmp_path,
        )
    )

    with pytest.raises(ValueError, match="not allowed"):
        loader.load_url("https://example.com/requirements")


def test_source_loader_rejects_allowlisted_host_resolving_private(
    tmp_path, monkeypatch
):
    loader = SourceLoader(
        Settings(
            allowed_source_hosts=("requirements.example",),
            reference_dir=tmp_path,
            workspace_dir=tmp_path,
            output_dir=tmp_path,
        )
    )
    monkeypatch.setattr(
        "requirements_alchemist.sources.socket.getaddrinfo",
        lambda *_args, **_kwargs: [
            (2, 1, 6, "", ("127.0.0.1", 443))
        ],
    )

    with pytest.raises(ValueError, match="Private or reserved"):
        loader.load_url("https://requirements.example/prd")


def test_jira_publisher_is_disabled_by_default():
    with pytest.raises(PermissionError, match="disabled"):
        JiraPublisher(Settings()).publish(_package("APPROVED"))


def test_chat_masks_configured_secrets_before_llm(monkeypatch):
    monkeypatch.setenv("TEST_RA_KEY", "never-send-this-value")
    llm = RequirementsLLM(Settings(llm_api_key_env="TEST_RA_KEY"))
    captured = {}

    def fake_chat(messages, **_kwargs):
        captured["messages"] = messages
        return "grounded answer"

    monkeypatch.setattr(llm, "_chat", fake_chat)
    llm.chat(
        "Does never-send-this-value appear?",
        [
            RetrievedChunk(
                id="source-1#chunk-1",
                source_id="source-1",
                title="PRD",
                text="Credential never-send-this-value must not leave the process.",
                score=1.0,
                metadata={},
            )
        ],
    )

    assert "never-send-this-value" not in json.dumps(captured["messages"])


def test_replay_web_workflow_generates_and_approves_story(tmp_path):
    reference_dir = tmp_path / "references"
    reference_dir.mkdir()
    settings = Settings(
        llm_provider="replay",
        reference_dir=reference_dir,
        workspace_dir=tmp_path / "workspace",
        output_dir=tmp_path / "output",
    )
    app = create_app(settings)
    app.testing = True
    client = app.test_client()
    denied = client.post(
        "/api/sources/text",
        json={"title": "No CSRF", "text": "Must not be accepted."},
    )
    client.get("/")
    with client.session_transaction() as session:
        csrf = session["csrf"]
    headers = {"X-CSRF-Token": csrf}

    source = client.post(
        "/api/sources/text",
        json={
            "title": "Account PRD",
            "text": "Registered users must sign in before viewing their private workspace.",
        },
        headers=headers,
    )
    generated = client.post(
        "/api/generate",
        json=GenerationInput(
            title="Account backlog", objective="Secure account access"
        ).model_dump(),
        headers=headers,
    )
    reviewed = client.patch(
        "/api/stories/STORY-001/review",
        json={"status": "APPROVED", "reviewer": "Vivek"},
        headers=headers,
    )

    assert denied.status_code == 403
    assert source.status_code == 201
    assert generated.status_code == 201
    assert generated.get_json()["stories"][0]["review_status"] == "DRAFT"
    assert reviewed.get_json()["review_status"] == "APPROVED"
    assert len(list(settings.output_dir.glob("pkg-*.json"))) == 1


def test_guided_demo_loads_three_detailed_draft_stories(tmp_path):
    settings = Settings(
        llm_provider="replay",
        reference_dir=tmp_path / "references",
        workspace_dir=tmp_path / "workspace",
        output_dir=tmp_path / "output",
    )
    settings.reference_dir.mkdir()
    app = create_app(settings)
    app.testing = True
    client = app.test_client()
    client.get("/")
    with client.session_transaction() as session:
        headers = {"X-CSRF-Token": session["csrf"]}

    response = client.post("/api/demo/load", headers=headers, json={})
    payload = response.get_json()

    assert response.status_code == 200
    assert payload["demo"] is True
    assert len(payload["package"]["stories"]) == 3
    assert all(
        story["review_status"] == "DRAFT"
        for story in payload["package"]["stories"]
    )
    assert payload["package"]["stories"][0]["acceptance_criteria"][0]["source_ids"] == [
        "demo-expense-prd"
    ]
