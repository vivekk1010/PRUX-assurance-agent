from pathlib import Path

import pytest
from openpyxl import load_workbook

from agent.excel_io import export_plan, import_review
from agent.models import Step, TestCase as PlanCase, TestPlan as Plan


def plan():
    return Plan(
        id="BLOG-1-v1-abc", version=1, story_key="BLOG-1",
        cases=[PlanCase(
            id="TC-1", story_key="BLOG-1", ac_ids=["AC-1"], title="Login",
            steps=[Step(action="goto", path="/login")],
        )],
    )


def test_excel_review_round_trip_with_step_edit(tmp_path: Path):
    source = plan()
    path = export_plan(source, tmp_path / "review.xlsx", secrets=["secret"])
    book = load_workbook(path)
    review = book["Human Review"]
    review["B2"], review["C2"], review["D2"] = "APPROVED", "Reviewer", "Looks good"
    book["Steps"]["K2"] = "/signin"
    book.save(path)

    reviewed = import_review(
        path, source, base_url="https://app.example.com",
        secrets=["secret"], valid_ac_ids={"AC-1"},
    )
    assert reviewed.cases[0].status == "APPROVED"
    assert reviewed.cases[0].reviewer == "Reviewer"
    assert reviewed.cases[0].steps[0].path == "/signin"


def test_excel_rejects_formula_and_secret(tmp_path: Path):
    path = export_plan(plan(), tmp_path / "review.xlsx")
    book = load_workbook(path)
    book["Human Review"]["D2"] = "=HYPERLINK(\"bad\")"
    book.save(path)
    with pytest.raises(ValueError, match="formula"):
        import_review(path, plan(), base_url="https://app.example.com")

    path = export_plan(plan(), tmp_path / "secret.xlsx")
    book = load_workbook(path)
    book["Human Review"]["D2"] = "hunter2"
    book.save(path)
    with pytest.raises(ValueError, match="secret"):
        import_review(path, plan(), base_url="https://app.example.com", secrets=["hunter2"])


def test_excel_rejects_external_navigation(tmp_path: Path):
    path = export_plan(plan(), tmp_path / "review.xlsx")
    book = load_workbook(path)
    book["Steps"]["K2"] = "https://evil.example/"
    book.save(path)
    with pytest.raises(ValueError, match="refused"):
        import_review(path, plan(), base_url="https://app.example.com")
